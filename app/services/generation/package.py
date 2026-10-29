"""Paquete ZIP del borrador (F3): DOCX, matriz CSV/XLSX, README de pendientes y manifiesto con huellas.

`build_package(snapshot, rendered)` es una función pura: mismas entradas, mismos bytes (las entradas del ZIP llevan la
fecha de generación de la instantánea). No empaqueta archivos de documentos: los del expediente solo se listan en
`excluded_documents` con su motivo; los de terceros nunca se redistribuyen.
"""
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core import config
from app.services.generation import formatting as fmt
from app.services.generation.docx_templates import DRAFT_NOTICE, HEADER_TEXT, TEMPLATE_FILES, TEMPLATE_VERSION
from app.services.generation.formatting import pending
from app.services.generation.matrix import COLUMNS, matrix_rows
from app.services.generation.ziputil import normalize_zip, zip_date

DATA_FILES = ("matriz.csv", "matriz.xlsx", "README_PENDIENTES.md")
KINDS = ("REAL", "SYNTHETIC", "HISTORICAL")


# ── Seguridad de celdas ──────────────────────────────────────────────────────

def safe_cell(value) -> str:
    """Texto seguro para CSV/XLSX: sin caracteres de control XML y sin que una hoja de cálculo lo ejecute."""
    text = fmt.clean_text(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


# ── CSV ──────────────────────────────────────────────────────────────────────

def build_csv(snapshot: dict) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(COLUMNS + ["Aviso"])
    for row in matrix_rows(snapshot):
        writer.writerow([safe_cell(c) for c in row] + [snapshot["notice"]])
    return buf.getvalue().encode("utf-8-sig")


# ── XLSX ─────────────────────────────────────────────────────────────────────

_HEAD_FILL = PatternFill("solid", fgColor="0F75BC")
_NOTICE_FILL = PatternFill("solid", fgColor="FFF4D6")


def _sheet_header(ws, snapshot: dict, columns: list[str], widths: list[int]) -> None:
    ws["A1"] = snapshot["notice"]
    ws["A1"].font = Font(bold=True, color="001D4C")
    ws["A1"].fill = _NOTICE_FILL
    ws["A2"] = DRAFT_NOTICE
    ws["A2"].font = Font(italic=True)
    for i, (title, width) in enumerate(zip(columns, widths), 1):
        cell = ws.cell(row=3, column=i, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A4"
    ws.oddHeader.center.text = HEADER_TEXT
    ws.oddFooter.center.text = snapshot["notice"]


def build_xlsx(snapshot: dict) -> bytes:
    wb = Workbook()
    generated = datetime.fromisoformat(snapshot["generated_at"]).astimezone(timezone.utc).replace(tzinfo=None)
    wb.properties.creator = "Dipromedic · Demo"
    wb.properties.title = f"Matriz de cumplimiento · {snapshot['opportunity']['external_id']}"
    wb.properties.created = wb.properties.modified = generated

    # Resumen
    ws = wb.active
    ws.title = "Resumen"
    _sheet_header(ws, snapshot, ["Concepto", "Valor"], [34, 110])
    opp = snapshot["opportunity"]
    totals = snapshot["totals"]
    summary = [
        ("Oportunidad", opp["external_id"]), ("Título", opp["title"]),
        ("Procedencia del dato", fmt.label(fmt.DATA_KIND_LABEL, opp["data_kind"])),
        ("Versión de datos", snapshot["data_version"]), ("Versión de plantillas", TEMPLATE_VERSION),
        ("Generado", fmt.stamp(snapshot["generated_at"])),
        ("Requisitos abiertos", snapshot["open_requirements"]), ("Total", totals["text"]),
        ("Tratamiento fiscal", snapshot["fiscal"]["text"]),
    ] + [("Advertencia", w) for w in snapshot["warnings"]]
    for r, (key, value) in enumerate(summary, 4):
        ws.cell(row=r, column=1, value=key).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=value if isinstance(value, int) else safe_cell(value))
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Matriz
    ws = wb.create_sheet("Matriz")
    _sheet_header(ws, snapshot, COLUMNS, [16, 11, 46, 13, 13, 22, 36, 44, 44, 40, 24, 22, 20])
    for r, row in enumerate(matrix_rows(snapshot), 4):
        for c, value in enumerate(row, 1):
            ws.cell(row=r, column=c, value=safe_cell(value)).alignment = Alignment(wrap_text=True, vertical="top")

    # Partidas: el servidor calcula con Decimal y escribe números de hoja de cálculo, que Excel guarda en coma
    # flotante (a estas magnitudes ida y vuelta da el mismo Decimal). El valor exacto de referencia es el texto de
    # los DOCX y de la instantánea.
    ws = wb.create_sheet("Partidas")
    _sheet_header(ws, snapshot, ["Partida", "Descripción", "Cantidad", "Unidad", "Precio unitario", "Moneda",
                                "Subtotal", "Origen del precio", "Variante propuesta"],
                  [10, 52, 12, 12, 16, 10, 18, 28, 48])
    row = 3
    for lot in snapshot["lots"]:
        row += 1
        has_price = lot["unit_price"] is not None
        origin = fmt.label(fmt.PRICE_KIND_LABEL, lot["price_data_kind"]) if lot["price_data_kind"] else \
            (pending("procedencia del precio") if has_price else "—")
        if lot["price_data_kind"] == "SYNTHETIC" and lot["price_source"] == "variante":
            origin += " (catálogo ficticio)"
        variant = lot["variant"]
        ws.cell(row=row, column=1, value=safe_cell(lot["number"]))
        ws.cell(row=row, column=2, value=safe_cell(lot["description"])).alignment = Alignment(wrap_text=True)
        quantity = Decimal(lot["quantity"]) if lot["quantity"] else None
        qty = ws.cell(row=row, column=3, value=quantity if quantity is not None else pending("cantidad"))
        # Entero sin punto final («40», no «40.»); decimales solo cuando la cantidad los tiene.
        qty.number_format = "#,##0" if quantity is not None and quantity == quantity.to_integral_value() \
            else "#,##0.###"
        ws.cell(row=row, column=4, value=safe_cell(lot["unit"] or pending("unidad")))
        price = ws.cell(row=row, column=5, value=Decimal(lot["unit_price"]) if has_price else pending("precio"))
        price.number_format = "#,##0.00"
        ws.cell(row=row, column=6, value=(lot["currency"] or pending("moneda")) if has_price else "—")
        sub = ws.cell(row=row, column=7,
                      value=Decimal(lot["subtotal"]) if lot["subtotal"] is not None else pending("precio"))
        sub.number_format = "#,##0.00"
        ws.cell(row=row, column=8, value=origin)
        ws.cell(row=row, column=9, value=safe_cell(f"{variant['sku']} · {variant['name']}" if variant
                                                   else pending("variante propuesta")))
    row += 2
    ws.cell(row=row, column=2, value=totals["text"]).font = Font(bold=True)
    ws.cell(row=row + 1, column=2, value=snapshot["fiscal"]["text"])
    ws.cell(row=row + 2, column=2, value="Los montos de esta hoja son números de hoja de cálculo (coma flotante); "
                                         "el valor exacto de referencia es el texto de los documentos DOCX.")

    buf = io.BytesIO()
    wb.save(buf)
    # openpyxl sella la fecha de modificación con el reloj al guardar: se fija la de la instantánea.
    stamp = generated.strftime("%Y-%m-%dT%H:%M:%SZ").encode()

    def fix_core(xml: bytes) -> bytes:
        return re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, xml)

    return normalize_zip(buf.getvalue(), zip_date(snapshot["generated_at"]), {"docProps/core.xml": fix_core})


# ── README de pendientes ─────────────────────────────────────────────────────

def _md(value) -> str:
    return fmt.clean_text(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ")


def excluded_documents(snapshot: dict) -> list[dict]:
    out = []
    for d in snapshot["documents"]:
        if d["third_party"]:
            reason = "Documento de tercero: solo se enlaza, no se redistribuye"
        elif d["has_file"]:
            reason = "No se incluye el archivo: este borrador no empaqueta documentos del Centro documental"
        else:
            reason = "Sin archivo adjunto en la plataforma"
        out.append({"code": d["code"], "title": d["title"], "data_kind": d["data_kind"],
                    "third_party": d["third_party"], "external_url": d["external_url"] if d["third_party"] else None,
                    "reason": reason})
    return out


def build_readme(snapshot: dict, excluded: list[dict]) -> str:
    opp = snapshot["opportunity"]
    open_reqs = [r for r in snapshot["requirements"] if r["open"]]
    lines = [
        "# Pendientes del borrador de expediente", "",
        f"> {snapshot['notice']}", f"> {DRAFT_NOTICE}", "",
        f"- Oportunidad: {_md(opp['external_id'])} · {_md(opp['title'])}",
        f"- Procedencia del dato: {fmt.label(fmt.DATA_KIND_LABEL, opp['data_kind'])}",
        f"- Versión de datos: `{snapshot['data_version']}`",
        f"- Versión de plantillas: {TEMPLATE_VERSION}",
        f"- Generado: {fmt.stamp(snapshot['generated_at'])}", "",
        f"## Requisitos abiertos ({len(open_reqs)})", "",
    ]
    if open_reqs:
        lines += ["| Código | Partida | Requisito | Estado | Responsable | Fecha límite | Acción siguiente |",
                  "|---|---|---|---|---|---|---|"]
        for r in open_reqs:
            task = r["open_task"] or {}
            lines.append("| " + " | ".join(_md(x) for x in [
                r["code"] or "s/c", f"Partida {r['lot']}" if r["lot"] else "General", r["text"],
                fmt.label(fmt.STATUS_LABEL, r["status"]),
                task.get("assignee") or r["reviewer"] or pending("responsable"),
                fmt.long_date(task.get("due_date") or r["due_date"]) or pending("fecha límite"),
                (r["next_action"] or (f"Tarea {task['code']}" if task else None)) or pending("acción siguiente"),
            ]) + " |")
    elif snapshot["requirements"]:
        lines.append("No hay requisitos abiertos en este momento.")
    else:
        lines.append("No hay requisitos capturados: la matriz está vacía. Eso no implica cumplimiento.")
    lines += ["", "## Advertencias", ""] + [f"- {_md(w)}" for w in snapshot["warnings"]]
    lines += ["", "## Precios y totales", "", f"- {snapshot['totals']['text']}", f"- {snapshot['fiscal']['text']}"]
    lines += ["", "## Documentos que no se incluyen en el paquete", ""]
    lines += [f"- {_md(d['code'] or 's/c')} · {_md(d['title'])}: {d['reason']}" for d in excluded] or \
             ["- Ninguno vinculado al procedimiento."]
    lines += ["", "## Cómo verificar el paquete", "",
              "`manifest.json` lista cada archivo con su huella SHA-256 y su tamaño; el propio manifiesto no se lista "
              "a sí mismo. La huella del ZIP completo queda en el historial de generaciones de la plataforma.", ""]
    return "\n".join(lines)


# ── Manifiesto y ZIP ─────────────────────────────────────────────────────────

def _provenance(snapshot: dict) -> tuple[dict, dict]:
    def count(kinds) -> dict:
        out: dict[str, int] = {}
        for kind in kinds:
            if kind:
                out[kind] = out.get(kind, 0) + 1
        return dict(sorted(out.items()))

    detail = {
        "opportunity": {snapshot["opportunity"]["data_kind"]: 1},
        "requirements": count(r["data_kind"] for r in snapshot["requirements"]),
        "variants": count(l["variant"]["data_kind"] for l in snapshot["lots"] if l["variant"]),
        "prices": count(l["price_data_kind"] for l in snapshot["lots"] if l["unit_price"] is not None),
        "documents": count(d["data_kind"] for d in snapshot["documents"]),
    }
    total = {k: sum(part.get(k, 0) for part in detail.values()) for k in KINDS}
    return total, detail


def build_package(snapshot: dict, rendered: dict[str, bytes]) -> tuple[bytes, dict]:
    """ZIP con `docx/*.docx`, `matriz.csv`, `matriz.xlsx`, `README_PENDIENTES.md` y `manifest.json`; y su manifiesto."""
    files: dict[str, bytes] = {}
    for key in TEMPLATE_FILES:
        if key in rendered:
            files[f"docx/{TEMPLATE_FILES[key][0]}"] = rendered[key]
    for key in rendered:
        if key not in TEMPLATE_FILES:
            files[f"docx/{key}.docx"] = rendered[key]
    excluded = excluded_documents(snapshot)
    files["matriz.csv"] = build_csv(snapshot)
    files["matriz.xlsx"] = build_xlsx(snapshot)
    files["README_PENDIENTES.md"] = build_readme(snapshot, excluded).encode("utf-8")

    provenance, detail = _provenance(snapshot)
    manifest = {
        "files": [{"name": name, "sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data)}
                  for name, data in files.items()],
        "template_version": TEMPLATE_VERSION,
        "generated_at": snapshot["generated_at"],
        "data_version": snapshot["data_version"],
        "opportunity": {"external_id": snapshot["opportunity"]["external_id"],
                        "data_kind": snapshot["opportunity"]["data_kind"]},
        "provenance": provenance,
        "provenance_detail": detail,
        "warnings": list(snapshot["warnings"]),
        "notice": config.WATERMARK,
        "excluded_documents": excluded,
    }
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")

    date_time = zip_date(snapshot["generated_at"])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            info = zipfile.ZipInfo(name, date_time=date_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, data)
    return buf.getvalue(), manifest
