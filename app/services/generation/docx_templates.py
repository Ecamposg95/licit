"""Seis plantillas DOCX P0 (F2).

Cada `render_*` recibe la instantánea de F1 y devuelve los bytes de un DOCX editable con:
- encabezado «Dipromedic · Borrador de expediente · DEMO»;
- pie con la marca DEMO literal, la versión de datos (`data_version`) y la fecha de generación;
- cada dato faltante como `[PENDIENTE: …]` (resaltado en amarillo), nunca un valor inventado;
- sin espacios para rúbricas ni textos que den el borrador por entregado.
Los textos que vienen de los datos (requisitos, títulos) se citan tal cual; el texto propio de las plantillas
no usa las palabras reservadas del brief.
"""
import io
import re
from datetime import datetime, timezone
from typing import Callable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.services.generation import formatting as fmt
from app.services.generation.formatting import pending
from app.services.generation.matrix import (
    assessment_lines, evidence_lines, mandatory_text, next_action, requirement_text,
)
from app.services.generation.ziputil import normalize_zip, zip_date

TEMPLATE_VERSION = "1.0"
HEADER_TEXT = "Dipromedic · Borrador de expediente · DEMO"
DRAFT_NOTICE = "Borrador · no apto para presentación"

# Archivo dentro del ZIP y título de cada documento, en el orden del expediente.
TEMPLATE_FILES = {
    "resumen": ("01_resumen_oportunidad.docx", "Resumen de oportunidad"),
    "matriz": ("02_matriz_cumplimiento.docx", "Matriz de cumplimiento"),
    "faltantes": ("03_lista_faltantes.docx", "Lista de faltantes con responsables"),
    "indice": ("04_indice_expediente.docx", "Índice de expediente"),
    "propuesta_tecnica": ("05_propuesta_tecnica_preliminar.docx", "Propuesta técnica preliminar"),
    "propuesta_economica": ("06_propuesta_economica_demo.docx", "Propuesta económica de demostración"),
}

NAVY = RGBColor(0x00, 0x1D, 0x4C)
BLUE_HEX = "0F75BC"
LIGHT_HEX = "E8F1F9"
AMBER_HEX = "FFF4D6"
_PENDING = re.compile(r"(\[PENDIENTE:[^\]]*\])")


# ── Construcción de bajo nivel ───────────────────────────────────────────────

def _shade(cell, hex_fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def _runs(paragraph, text, *, bold=False, italic=False, size=None, color=None) -> None:
    """Agrega `text` al párrafo resaltando en amarillo cada marcador [PENDIENTE: …]."""
    for chunk in _PENDING.split(fmt.clean_text(text)):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        run.bold = bold or None
        run.italic = italic or None
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if _PENDING.fullmatch(chunk):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True


def _para(doc, text="", *, bold=False, italic=False, size=None, style=None, align=None, color=None, space_after=4):
    p = doc.add_paragraph(style=style)
    if text:
        _runs(p, text, bold=bold, italic=italic, size=size, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def _heading(doc, text: str, level: int = 1):
    h = doc.add_heading(level=level)
    _runs(h, text, color=NAVY)
    return h


def _bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    _runs(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def _banner(doc, text: str, fill: str = AMBER_HEX):
    """Aviso sombreado (párrafo, no tabla) para que las tablas de datos sigan siendo las únicas del documento."""
    p = doc.add_paragraph()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)
    _runs(p, text, bold=True, size=10)
    p.paragraph_format.space_after = Pt(6)
    return p


def _write_cell(cell, value) -> None:
    """Un valor de celda puede ser texto (con saltos de línea) o una lista de textos."""
    lines = value if isinstance(value, list) else str(value if value is not None else "").split("\n")
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(1)
        _runs(p, line, size=9)


def _table(doc, headers: list[str], rows: list[list], widths_cm: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    head = table.rows[0]
    trpr = head._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for cell, title in zip(head.cells, headers):
        _shade(cell, BLUE_HEX)
        cell.paragraphs[0].text = ""
        _runs(cell.paragraphs[0], title, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            _write_cell(cell, value)
    if widths_cm:
        for row in table.rows:
            for cell, width in zip(row.cells, widths_cm):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _kv_table(doc, pairs: list[tuple[str, str]], widths_cm=(5.5, 11.5)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for key, value in pairs:
        cells = table.add_row().cells
        _shade(cells[0], LIGHT_HEX)
        cells[0].paragraphs[0].text = ""
        _runs(cells[0].paragraphs[0], key, bold=True, size=9)
        _write_cell(cells[1], value)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _new_document(snapshot: dict, title: str, *, landscape: bool = False):
    doc = Document()
    generated = datetime.fromisoformat(snapshot["generated_at"])
    if generated.tzinfo is None:
        generated = generated.replace(tzinfo=timezone.utc)
    props = doc.core_properties
    props.title = f"{title} · {snapshot['opportunity']['external_id']}"
    props.author = "Dipromedic · Demo"
    props.last_modified_by = "Dipromedic · Demo"
    props.subject = DRAFT_NOTICE
    props.comments = snapshot["notice"]
    props.created = props.modified = generated

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10)

    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(1.8))
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)

    header = section.header.paragraphs[0]
    header.text = ""
    _runs(header, HEADER_TEXT, bold=True, size=9, color=NAVY)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(footer, snapshot["notice"], bold=True, size=8, color=NAVY)
    detail = section.footer.add_paragraph()
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(detail, f"Versión de datos: {snapshot['data_version']} · Generado: {fmt.stamp(snapshot['generated_at'])}",
          size=7)

    _heading(doc, title, 0)
    opp = snapshot["opportunity"]
    _para(doc, f"{opp['external_id']} · {opp['title']}", bold=True, size=11, space_after=6)
    _banner(doc, DRAFT_NOTICE)
    if opp["data_kind"] == "SYNTHETIC":
        _para(doc, "Escenario ficticio: este procedimiento no existe en ninguna fuente oficial.", italic=True, size=9)
    return doc


def _bytes(doc) -> bytes:
    """DOCX reproducible: mismas instantánea y plantilla, mismos bytes (entradas con la fecha de generación)."""
    buf = io.BytesIO()
    doc.save(buf)
    return normalize_zip(buf.getvalue(), zip_date(doc.core_properties.created))


# ── Datos compartidos ────────────────────────────────────────────────────────

def _or_pending(value, what: str) -> str:
    return str(value) if value not in (None, "", []) else pending(what)


def _open(snapshot: dict) -> list[dict]:
    return [r for r in snapshot["requirements"] if r["open"]]


def _lot_label(number) -> str:
    return f"Partida {number}" if number else "General"


def _affinity_text(aff: dict) -> str:
    if aff["state"] == "no_calculada":
        return "Afinidad no calculada"
    components = ", ".join(f"{k}: {v}" for k, v in aff["components"].items()) or "sin componentes"
    if aff["total"] is None:
        text = f"Afinidad con cobertura incompleta (sin total). Componentes: {components}."
    else:
        text = f"Afinidad {aff['total']} de 100. Componentes: {components}."
    if aff["blocking"]:
        text += f" Bloqueos: {', '.join(aff['blocking'])}."
    return text


def _sorted_open(snapshot: dict) -> list[dict]:
    def key(r):
        due = (r["open_task"] or {}).get("due_date") or r["due_date"] or "9999-12-31"
        return (0 if r["mandatory"] is True else 1 if r["mandatory"] is None else 2, due, r["code"] or "", r["id"])
    return sorted(_open(snapshot), key=key)


# ── 1. Resumen de oportunidad ────────────────────────────────────────────────

def render_resumen(snapshot: dict) -> bytes:
    opp = snapshot["opportunity"]
    doc = _new_document(snapshot, TEMPLATE_FILES["resumen"][1])
    precision = opp["submission_precision"]
    submission = fmt.instant(opp["submission_at"], precision)
    if submission and precision == "date":
        submission += " (la fuente solo publica el día)"
    source = opp["source"]["name"] if opp["source"] else None
    if source and opp["source_url"]:
        source += f" · {opp['source_url']}"
    _heading(doc, "Datos del procedimiento", 1)
    _kv_table(doc, [
        ("Identificador", opp["external_id"]),
        ("Título", opp["title"]),
        ("Convocante", opp["buyer"]),
        ("Unidad compradora", _or_pending(opp["buying_unit"], "unidad compradora")),
        ("Entidad federativa", _or_pending(opp["state"], "entidad federativa")),
        ("Tipo de procedimiento", _or_pending(fmt.label(fmt.PROCEDURE_LABEL, opp["procedure_type"]),
                                              "tipo de procedimiento")),
        ("Carácter", _or_pending(fmt.label(fmt.CHARACTER_LABEL, opp["character"]), "carácter del procedimiento")),
        ("Familias", _or_pending(", ".join(opp["families"]), "familias de producto")),
        ("Procedencia del dato", fmt.label(fmt.DATA_KIND_LABEL, opp["data_kind"])),
        ("Fuente", _or_pending(source, "fuente")),
        ("Estado reportado por la fuente", _or_pending(opp["reported_status"], "estado reportado")),
        ("Verificación", _or_pending(fmt.label(fmt.VERIFICATION_LABEL, opp["verification_status"]),
                                     "estado de verificación")
         + (f" · {opp['verification_scope']}" if opp["verification_scope"] else "")),
        ("Última comprobación", _or_pending(fmt.instant(opp["last_checked_at"]), "última comprobación")),
        ("Junta de aclaraciones", _or_pending(fmt.instant(opp["questions_at"]), "fecha de junta de aclaraciones")),
        ("Presentación y apertura", _or_pending(submission, "fecha de presentación")),
        ("Fallo", _or_pending(fmt.instant(opp["ruling_at"]), "fecha de fallo")),
        ("Monto estimado", _or_pending(fmt.money(opp["estimated_amount"], opp["currency"]), "monto estimado")),
        ("Responsable del procedimiento", _or_pending(opp["owner"], "responsable")),
        ("Revisión humana", "Requerida" if opp["requires_human_review"] else "No requerida"),
    ])

    _heading(doc, "Partidas", 1)
    if snapshot["lots"]:
        rows = []
        for lot in snapshot["lots"]:
            v = lot["variant"]
            rows.append([lot["number"], lot["description"],
                         _or_pending(fmt.quantity(lot["quantity"]), "cantidad"),
                         _or_pending(lot["unit"], "unidad"),
                         f"{v['sku']} · {v['name']}" if v else pending("variante propuesta")])
        _table(doc, ["Partida", "Descripción", "Cantidad", "Unidad", "Variante propuesta"], rows,
               [1.8, 7.5, 2.2, 2.2, 5.0])
        _para(doc, snapshot["totals"]["text"], bold=True)
    else:
        _para(doc, pending("partidas del procedimiento") + " No se infieren partidas que la fuente no publica.")

    _heading(doc, "Decisión y afinidad", 1)
    ws = snapshot["workspace"]
    _para(doc, "Decisión de participación: " + (fmt.label(fmt.DECISION_LABEL, ws["decision"]) if ws else
                                                 "Sin decisión registrada"))
    _para(doc, _affinity_text(snapshot["affinity"]))

    _heading(doc, "Estado de los requisitos", 1)
    reqs = snapshot["requirements"]
    if reqs:
        counts: dict[str, int] = {}
        for r in reqs:
            counts[r["status"]] = counts.get(r["status"], 0) + 1
        _table(doc, ["Estado", "Requisitos"], [[fmt.label(fmt.STATUS_LABEL, s), str(n)]
                                                for s, n in sorted(counts.items())], [6, 3])
        _para(doc, f"Requisitos abiertos: {snapshot['open_requirements']} de {len(reqs)}. Ver la lista de faltantes.")
    else:
        _para(doc, "No hay requisitos capturados: la matriz está vacía. Eso no implica cumplimiento.")

    _heading(doc, "Advertencias del borrador", 1)
    for warning in snapshot["warnings"]:
        _bullet(doc, warning)
    return _bytes(doc)


# ── 2. Matriz de cumplimiento ────────────────────────────────────────────────

def render_matriz(snapshot: dict) -> bytes:
    doc = _new_document(snapshot, TEMPLATE_FILES["matriz"][1], landscape=True)
    reqs = snapshot["requirements"]
    _para(doc, "El estado de cada requisito lo fija una persona con permiso de revisión. Una evidencia vinculada o "
               "una evaluación automática no concede cumplimiento por sí solas.", italic=True, size=9)
    if not reqs:
        _para(doc, "No hay requisitos capturados: la matriz está vacía. Eso no implica cumplimiento.", bold=True)
        return _bytes(doc)
    rows = []
    for r in reqs:
        status = [fmt.label(fmt.STATUS_LABEL, r["status"])]
        if r["status_reason"]:
            status.append(r["status_reason"])
        if r["open"]:
            status.append(pending("validación del requisito"))
        evidence = evidence_lines(r)
        if not evidence:
            evidence = [pending("evidencia vinculada")] if r["open"] else ["—"]
        rows.append([r["code"] or "s/c", _lot_label(r["lot"]), requirement_text(r),
                     fmt.label(fmt.REQ_TYPE_LABEL, r["req_type"], "—"), mandatory_text(r), status, evidence,
                     assessment_lines(r["assessments"]),
                     next_action(r) or "—"])
    _table(doc, ["Código", "Partida", "Requisito", "Tipo", "Obligatorio", "Estado", "Evidencia vinculada",
                 "Evaluación técnica automática", "Siguiente acción"], rows,
           [2.2, 1.8, 5.2, 1.9, 2.0, 3.6, 4.4, 4.0, 3.2])
    _para(doc, f"Requisitos abiertos: {snapshot['open_requirements']} de {len(reqs)}.", bold=True)
    return _bytes(doc)


# ── 3. Lista de faltantes con responsables ───────────────────────────────────

def render_faltantes(snapshot: dict) -> bytes:
    doc = _new_document(snapshot, TEMPLATE_FILES["faltantes"][1], landscape=True)
    open_reqs = _sorted_open(snapshot)
    _heading(doc, "Requisitos abiertos", 1)
    if open_reqs:
        rows = []
        for i, r in enumerate(open_reqs, 1):
            task = r["open_task"] or {}
            owner = task.get("assignee") or r["reviewer"]
            due = task.get("due_date") or r["due_date"]
            action = next_action(r)
            req_cell = [r["text"]]
            if r["status_reason"]:
                req_cell.append(f"Motivo: {r['status_reason']}")
            rows.append([str(i), r["code"] or "s/c", req_cell,
                         {True: "Obligatorio", False: "Opcional", None: "Por determinar"}[r["mandatory"]],
                         [fmt.label(fmt.STATUS_LABEL, r["status"]), pending("cierre del requisito")],
                         _or_pending(owner, "responsable"),
                         _or_pending(fmt.long_date(due), "fecha límite"),
                         _or_pending(action, "acción siguiente")])
        _table(doc, ["No.", "Código", "Requisito", "Carácter", "Estado", "Responsable", "Fecha límite",
                     "Acción siguiente"], rows, [1.1, 2.3, 6.8, 2.2, 2.6, 3.4, 3.0, 5.0])
    elif snapshot["requirements"]:
        _para(doc, "No hay requisitos abiertos en este momento. Cada cierre lo validó una persona con permiso.")
    else:
        _para(doc, "No hay requisitos capturados: la matriz está vacía. Eso no implica cumplimiento.", bold=True)

    _heading(doc, "Advertencias y otros pendientes del borrador", 1)
    others = [w for w in snapshot["warnings"] if "requisito(s) abierto(s)" not in w]
    for warning in others:
        _bullet(doc, warning)
    no_validity = [d for d in snapshot["documents"] if not d["valid_until"]]
    for d in no_validity:
        _bullet(doc, f"Documento {d['code'] or d['title']}: " + pending("vigencia"))
    return _bytes(doc)


# ── 4. Índice de expediente ──────────────────────────────────────────────────

def render_indice(snapshot: dict, included: list[str] | None = None) -> bytes:
    """`included`: claves de las plantillas que viajan en este paquete (omitido = las seis)."""
    doc = _new_document(snapshot, TEMPLATE_FILES["indice"][1])
    _heading(doc, "A. Documentos generados en este paquete", 1)
    rows = [[str(i), f"docx/{name}", title]
            for i, (name, title) in enumerate((v for k, v in TEMPLATE_FILES.items()
                                               if included is None or k in included), 1)]
    _table(doc, ["No.", "Archivo", "Contenido"], rows, [1.2, 8.0, 8.0])
    _heading(doc, "B. Archivos de datos", 1)
    _table(doc, ["Archivo", "Contenido"], [
        ["matriz.csv", "Matriz de cumplimiento en CSV (UTF-8 con BOM)"],
        ["matriz.xlsx", "Matriz de cumplimiento, partidas y resumen en hoja de cálculo"],
        ["README_PENDIENTES.md", "Pendientes y advertencias del borrador"],
        ["manifest.json", "Archivos del paquete con su huella SHA-256, versión de plantillas y de datos"],
    ], [5.0, 12.2])

    _heading(doc, "C. Documentos del expediente", 1)
    docs = snapshot["documents"]
    if docs:
        rows = []
        for d in docs:
            if d["third_party"]:
                link = f" ({d['external_url']})" if d["external_url"] else ""
                where = "Documento de tercero: solo se enlaza" + link
            elif d["has_file"]:
                where = "No se incluye el archivo en el borrador; se consulta en el Centro documental"
            else:
                where = "Sin archivo adjunto en la plataforma"
            rows.append([d["code"] or "s/c", d["title"], fmt.label(fmt.DATA_KIND_LABEL, d["data_kind"]),
                         _or_pending(fmt.long_date(d["valid_until"]), "vigencia"), d["review_status"], where])
        _table(doc, ["Código", "Título", "Procedencia", "Vigencia", "Revisión", "Archivo"], rows,
               [2.2, 5.4, 2.6, 3.0, 1.8, 4.2])
    else:
        _para(doc, pending("documentos del expediente") + " No hay documentos vinculados al procedimiento.")

    _heading(doc, "D. Documentos por conseguir", 1)
    missing = [r for r in _sorted_open(snapshot) if not r["evidence"]]
    if missing:
        for r in missing:
            _bullet(doc, pending(f"documento para {r['code'] or 's/c'}: {r['text']}"))
    else:
        _para(doc, "Todos los requisitos abiertos tienen al menos una evidencia vinculada; "
                   "vinculada sin revisar no equivale a validada.")
    return _bytes(doc)


# ── 5. Propuesta técnica preliminar ──────────────────────────────────────────

def _describe_attribute(raw, default_unit: str | None = None) -> tuple[str | None, str | None]:
    """(texto, cita) de un atributo de catálogo en cualquiera de sus formas; (None, None) si no está declarado."""
    if raw is None or raw == "":
        return None, None
    if isinstance(raw, dict):
        value = raw.get("values", raw.get("value"))
        if value is None or value == [] or value == "":
            return None, None
        if isinstance(value, list):
            value = ", ".join(str(v) for v in value)
        unit = raw.get("unit") or default_unit
        return f"{value}{' ' + unit if unit else ''}", raw.get("source_text")
    return f"{raw}{' ' + default_unit if default_unit and isinstance(raw, (int, float)) else ''}", None


def _spec_text(spec) -> str:
    if isinstance(spec, dict) and "value" in spec:
        op = fmt.label(fmt.OPERATOR_LABEL, spec.get("operator"), "igual a")
        unit = f" {spec['unit']}" if spec.get("unit") else ""
        return f"{op} {spec['value']}{unit}"
    return str(spec)


def render_propuesta_tecnica(snapshot: dict) -> bytes:
    doc = _new_document(snapshot, TEMPLATE_FILES["propuesta_tecnica"][1])
    _para(doc, "Documento preliminar. Cada dato del catálogo se cita con su fuente; nada de lo que aquí aparece "
               "concede cumplimiento. La evaluación técnica la hace el motor de correspondencia y la valida una "
               "persona con permiso de revisión.", italic=True, size=9)
    if not snapshot["lots"]:
        _para(doc, pending("partidas del procedimiento") + " No se proponen productos sin partidas.")
        return _bytes(doc)
    for lot in snapshot["lots"]:
        qty = fmt.quantity(lot["quantity"])
        _heading(doc, f"Partida {lot['number']} · {lot['description']}", 1)
        _para(doc, f"Cantidad: {qty or pending('cantidad')} {lot['unit'] or ''}".rstrip(), size=9)
        v = lot["variant"]
        if v is None:
            _para(doc, "Variante propuesta: " + pending("variante propuesta"))
        else:
            product = v["product"] or {}
            _kv_table(doc, [
                ("Variante propuesta", f"{v['sku']} · {v['name']}"),
                ("Marca / fabricante", f"{product.get('brand') or pending('marca')} / "
                                       f"{product.get('manufacturer') or pending('fabricante')}"),
                ("Procedencia del dato", fmt.label(fmt.DATA_KIND_LABEL, v["data_kind"])),
                ("Registro sanitario", v["regulatory_reference"] or pending("registro sanitario")),
            ])
            if v["catalog_credit"]:
                _para(doc, v["catalog_credit"], italic=True, size=8)
        reqs_by_attr = {r["attribute"]: r for r in snapshot["requirements"]
                        if r["lot"] == lot["number"] and r["attribute"]}
        rows = []
        for key, spec in lot["specs"].items():
            unit = spec.get("unit") if isinstance(spec, dict) else None
            declared, cite = _describe_attribute((v["attributes"] if v else {}).get(key), unit)
            declared_cell = [declared if declared is not None else pending("dato de catálogo")]
            if cite:
                declared_cell.append(f"Cita: «{cite}»")
            req = reqs_by_attr.get(key)
            rows.append([fmt.label(fmt.ATTRIBUTE_LABEL, key), _spec_text(spec), declared_cell,
                         assessment_lines(req["assessments"] if req else [])])
        if rows:
            _table(doc, ["Atributo", "Exigido por las bases", "Declarado en el catálogo", "Evaluación técnica"],
                   rows, [3.4, 4.2, 5.0, 4.8])
        else:
            _para(doc, pending("especificaciones exigidas por las bases"))
        for line in (a for a in lot["assessments"]):
            _bullet(doc, f"Evaluación de la partida: {line['result']}: {line['explanation']}")
        open_tech = [r for r in _open(snapshot) if r["lot"] == lot["number"] and r["req_type"] == "tecnico"]
        if open_tech:
            codes = ", ".join(r["code"] or "s/c" for r in open_tech)
            _para(doc, f"Requisitos técnicos abiertos de la partida: {codes}.", size=9)
    _heading(doc, "Alcance", 1)
    _para(doc, "Suministrar una pieza no cubre por sí mismo los servicios o el instrumental que pidan las bases: "
               "esos requisitos se revisan en la matriz de cumplimiento.", size=9)
    return _bytes(doc)


# ── 6. Propuesta económica de demostración ───────────────────────────────────

def render_propuesta_economica(snapshot: dict) -> bytes:
    doc = _new_document(snapshot, TEMPLATE_FILES["propuesta_economica"][1], landscape=True)
    totals = snapshot["totals"]
    if totals["has_synthetic_prices"]:
        _banner(doc, "Los precios marcados SIMULADO salen del catálogo ficticio; no constituyen una cotización.")
    rows = []
    for lot in snapshot["lots"]:
        has_price = lot["unit_price"] is not None
        currency = lot["currency"]
        origin = "—"
        if has_price:
            origin = fmt.label(fmt.PRICE_KIND_LABEL, lot["price_data_kind"]) if lot["price_data_kind"] \
                else pending("procedencia del precio")
            if lot["price_data_kind"] == "SYNTHETIC" and lot["price_source"] == "variante":
                origin += "\n(catálogo ficticio)"
        qty = fmt.quantity(lot["quantity"])
        rows.append([
            lot["number"], lot["description"], qty or pending("cantidad"), lot["unit"] or pending("unidad"),
            fmt.money(lot["unit_price"], None) if has_price else pending("precio"),
            (currency or pending("moneda")) if has_price else "—",
            fmt.money(lot["subtotal"], None) if lot["subtotal"] is not None else pending("precio"),
            origin,
        ])
    if rows:
        _table(doc, ["Partida", "Descripción", "Cantidad", "Unidad", "Precio unitario", "Moneda", "Subtotal",
                     "Origen del precio"], rows, [1.8, 8.0, 2.0, 2.0, 3.0, 1.8, 3.2, 3.8])
    else:
        _para(doc, pending("partidas con precio") + " El procedimiento no tiene partidas capturadas.")
    _para(doc, totals["text"], bold=True, size=12)    # el texto lo compone el backend con Decimal
    _para(doc, "Descuentos: no se capturaron descuentos en la plataforma; no se aplicó ninguno.")
    _para(doc, snapshot["fiscal"]["text"], bold=True)
    ws = snapshot["workspace"] or {}
    approval = ws.get("commercial_approval")
    if approval:
        who = approval.get("by") or pending("persona que aprueba")
        _para(doc, f"Aprobación comercial de precios: {who}, {fmt.instant(approval['at']) or ''}"
              + (" (desactualizada por cambios posteriores)" if ws.get("approvals_stale") else ""))
    else:
        _para(doc, "Aprobación comercial de precios: " + pending("aprobación comercial de precios"))
    return _bytes(doc)


TEMPLATES: dict[str, Callable[[dict], bytes]] = {
    "resumen": render_resumen,
    "matriz": render_matriz,
    "faltantes": render_faltantes,
    "indice": render_indice,
    "propuesta_tecnica": render_propuesta_tecnica,
    "propuesta_economica": render_propuesta_economica,
}
