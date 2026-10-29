"""Frente F · generación de borradores: instantánea exacta, plantillas DOCX, paquete ZIP y trabajo asíncrono."""
import hashlib
import io
import json
import re
import time
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from sqlalchemy import select

from app.core import config
from app.core.database import SessionLocal
from app.models import (
    AuditEvent, BidWorkspace, DataKind, Document, GeneratedArtifact, Lot, MatchAssessment, Opportunity,
    Organization, Product, Requirement, Variant,
)
from app.services import audit
from app.services.generation.snapshot import FISCAL, _totals, build_snapshot, line_subtotal


# ── Escenario propio (no depende del re-sembrado de E) ───────────────────────

@pytest.fixture
def db():
    session = SessionLocal()
    yield session
    session.rollback()
    session.close()


def lic(db, code: str = "DEMO-LIC-001") -> Opportunity:
    return db.scalars(select(Opportunity).where(Opportunity.external_id == code)).one()


def set_lot_price(db, opp: Opportunity, number: str, price: str | None, kind: str | None = "SYNTHETIC",
                  currency: str | None = "MXN") -> Lot:
    lot = next(l for l in opp.lots if l.number == number)
    lot.unit_price = Decimal(price) if price is not None else None
    lot.currency = currency
    lot.price_data_kind = kind
    db.flush()
    return lot


def real_variant(db, org_id: str) -> Variant:
    """Variante REAL sin precio: el catálogo público no publica precios."""
    product = Product(organization_id=org_id, name="Placa real de prueba", family="Osteosíntesis",
                      data_kind="REAL", source_url="https://www.dipromedic.com/x", checked_on=date(2026, 10, 28))
    variant = Variant(organization_id=org_id, product=product, sku="DPM-TEST-1", name="Placa real 6 orificios",
                      unit="pieza", data_kind="REAL", attributes={"holes": {"values": [6], "unit": "orificios",
                                                                              "source_text": "6 orificios"}})
    db.add_all([product, variant])
    db.flush()
    return variant


def clear_lot_prices(db, opp: Opportunity, *numbers: str) -> None:
    """Quita el precio propio de las partidas indicadas (todas si no se indica): el precio sale de la variante."""
    for lot in opp.lots:
        if not numbers or lot.number in numbers:
            lot.unit_price = lot.currency = lot.price_data_kind = None
    db.flush()


def db_warnings(db) -> list[str]:
    return build_snapshot(db, lic(db))["warnings"]


def snapshot_of(db, code: str = "DEMO-LIC-001") -> dict:
    return build_snapshot(db, lic(db, code))


# ── F1 · instantánea y aritmética ────────────────────────────────────────────

def test_subtotal_exact():
    assert line_subtotal(Decimal("40"), Decimal("1500.00")) == Decimal("60000.00")
    assert str(line_subtotal(Decimal("40.000"), Decimal("1500.00"))) == "60000.00"
    # Redondeo comercial a centavos (mitad hacia arriba), sin errores de binario.
    assert line_subtotal(Decimal("3"), Decimal("0.125")) == Decimal("0.38")
    assert line_subtotal(Decimal("0.1"), Decimal("0.1")) == Decimal("0.01")
    assert line_subtotal(Decimal("3"), Decimal("33.335")) == Decimal("100.01")
    assert line_subtotal(None, Decimal("1")) is None
    assert line_subtotal(Decimal("1"), None) is None


def test_lot_price_beats_variant_price_and_is_exact(db):
    opp = lic(db)
    set_lot_price(db, opp, "1", "1500.00", "SYNTHETIC")
    snap = build_snapshot(db, opp)
    lot1 = next(l for l in snap["lots"] if l["number"] == "1")
    assert lot1["unit_price"] == "1500.00" and lot1["subtotal"] == "60000.00" and lot1["currency"] == "MXN"
    assert lot1["price_source"] == "partida"


def test_missing_price_warning_and_partial_total(db):
    opp = lic(db)
    # Partida 1: precio propio; partida 2: variante REAL (sin precio); partida 3: sin variante.
    set_lot_price(db, opp, "1", "1500.00", "SYNTHETIC")
    clear_lot_prices(db, opp, "2", "3")
    lot2 = next(l for l in opp.lots if l.number == "2")
    lot2.proposed_variant_id = real_variant(db, opp.organization_id).id
    lot3 = next(l for l in opp.lots if l.number == "3")
    lot3.proposed_variant_id = None
    db.flush()
    db.expire_all()
    snap = build_snapshot(db, lic(db))
    lots = {l["number"]: l for l in snap["lots"]}
    assert lots["1"]["subtotal"] == "60000.00"
    assert lots["2"]["unit_price"] is None and lots["2"]["subtotal"] is None
    assert lots["3"]["unit_price"] is None and lots["3"]["subtotal"] is None
    assert "Partida 2: precio pendiente" in snap["warnings"]
    assert "Partida 3: precio pendiente" in snap["warnings"]
    totals = snap["totals"]
    assert totals["known_total"] == "60000.00" and totals["currency"] == "MXN"
    assert totals["complete"] is False and totals["lots_priced"] == 1 and totals["lots_pending"] == 2
    assert totals["text"] == ("Total parcial: 60,000.00 MXN (solo 1 de 3 partidas con precio; "
                              "faltan las partidas 2 y 3)")
    assert snap["fiscal"] == FISCAL == {
        "status": "pendiente",
        "text": "Tratamiento fiscal pendiente de configurar; no se aplicaron impuestos."}


def test_partial_total_text_uses_singular_and_plural(db):
    opp = lic(db)
    clear_lot_prices(db, opp, "3")
    next(l for l in opp.lots if l.number == "3").proposed_variant_id = None
    db.flush()
    db.expire_all()
    text = build_snapshot(db, lic(db))["totals"]["text"]
    assert text == "Total parcial: 151,600.00 MXN (solo 2 de 3 partidas con precio; falta la partida 3)"
    # Tres o más: «2, 3 y 4».
    lots = [{"number": str(n), "subtotal": None if n > 1 else "10.00", "currency": "MXN", "unit_price": None,
             "price_data_kind": None} for n in range(1, 5)]
    assert _totals(lots)["text"].endswith("faltan las partidas 2, 3 y 4)")


def test_complete_total_sums_known_subtotals(db):
    snap = snapshot_of(db)    # precios ficticios de las partidas: 40×2450 + 20×2680 + 15×3120
    assert [l["subtotal"] for l in snap["lots"]] == ["98000.00", "53600.00", "46800.00"]
    assert snap["totals"]["known_total"] == "198400.00" and snap["totals"]["complete"] is True
    assert not [w for w in snap["warnings"] if "precio pendiente" in w]


def test_mixed_currencies_are_never_added_together(db):
    opp = lic(db)
    set_lot_price(db, opp, "1", "100.00", "SYNTHETIC", "MXN")
    set_lot_price(db, opp, "2", "10.00", "SYNTHETIC", "USD")
    db.expire_all()
    totals = build_snapshot(db, lic(db))["totals"]
    assert totals["currency"] is None and totals["known_total"] is None
    # Partida 1: 40×100; partida 3 (precio ficticio propio): 15×3120; partida 2 en dólares: 20×10.
    assert totals["by_currency"] == {"MXN": "50800.00", "USD": "200.00"}
    assert any("monedas" in w.lower() for w in db_warnings(db))


def test_data_version_stable(db):
    first = snapshot_of(db)
    second = snapshot_of(db)
    assert first["data_version"] == second["data_version"]
    assert re.fullmatch(r"[0-9a-f]{64}", first["data_version"])
    # Es el SHA-256 del JSON canónico sin generated_at.
    body = {k: v for k, v in first.items() if k not in ("generated_at", "data_version")}
    canonical = json.dumps(body, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    assert first["data_version"] == hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    # Un cambio real de datos cambia la versión.
    opp = lic(db)
    set_lot_price(db, opp, "1", "1600.00")
    db.expire_all()
    assert snapshot_of(db)["data_version"] != first["data_version"]


def test_synthetic_prices_flagged(db):
    snap = snapshot_of(db)
    assert snap["lots"], "el escenario debe tener partidas"
    for lot in snap["lots"]:
        assert lot["price_data_kind"] == "SYNTHETIC"
        assert lot["price_source"] == "partida"           # el escenario trae precios ficticios propios
    assert any("SIMULADO" in w for w in snap["warnings"])
    # Sin precio en la partida, el de la variante ficticia propuesta (también SIMULADO).
    clear_lot_prices(db, lic(db))
    db.expire_all()
    for lot in snapshot_of(db)["lots"]:
        assert lot["price_data_kind"] == "SYNTHETIC"
        assert lot["price_source"] == "variante"          # el precio viene del catálogo ficticio
        assert lot["variant"]["data_kind"] == "SYNTHETIC"
    snap = snapshot_of(db)
    assert snap["totals"]["has_synthetic_prices"] is True
    assert snap["notice"] == config.WATERMARK


def test_price_without_declared_origin_is_warned(db):
    opp = lic(db)
    set_lot_price(db, opp, "1", "1500.00", None)
    db.expire_all()
    snap = snapshot_of(db)
    lot1 = next(l for l in snap["lots"] if l["number"] == "1")
    assert lot1["price_data_kind"] is None
    assert any(w.startswith("Partida 1:") and "procedencia" in w for w in snap["warnings"])


def test_historical_reference_is_warned_and_has_no_lots(db):
    hist = db.scalars(select(Opportunity).where(Opportunity.data_kind == "HISTORICAL")).first()
    snap = build_snapshot(db, hist)
    assert "Referencia histórica: no es una oportunidad abierta" in snap["warnings"]
    assert snap["lots"] == [] and snap["totals"]["known_total"] is None
    assert "Total no calculable" in snap["totals"]["text"]
    body = docx_parts(TEMPLATES["propuesta_economica"](snap))["body"]
    assert "[PENDIENTE: partidas con precio]" in body


def test_open_requirements_evidence_and_task_in_snapshot(db):
    snap = snapshot_of(db)
    reqs = {r["code"]: r for r in snap["requirements"]}
    assert len(reqs) == len(snap["requirements"]) >= 6
    faltante = reqs["DEMO-REQ-005"]
    assert faltante["status"] == "faltante" and faltante["open"] is True and faltante["mandatory"] is True
    assert faltante["evidence"] == []
    assert snap["open_requirements"] == sum(1 for r in snap["requirements"] if r["open"])
    # Ningún requisito se cierra solo por tener evidencia.
    assert all(r["status"] != "cumple_validado" for r in snap["requirements"] if r["evidence"])


def test_affinity_is_opaque_and_null_means_not_computed(db):
    opp = lic(db)
    snap = build_snapshot(db, opp)
    assert snap["affinity"]["state"] == "no_calculada"
    empty = snap["workspace"]          # sin espacio creado: los valores de uno recién abierto
    assert empty["decision"] == "pendiente" and empty["technical_approval"] is None
    ws = BidWorkspace(organization_id=opp.organization_id, opportunity_id=opp.id,
                      affinity={"components": {"familia": 40}, "total": None, "coverage_complete": False,
                                "blocking": ["documental"]})
    db.add(ws)
    db.flush()
    snap = build_snapshot(db, opp)
    aff = snap["affinity"]
    assert aff["state"] == "calculada" and aff["total"] is None and aff["coverage_complete"] is False
    assert aff["blocking"] == ["documental"] and aff["components"] == {"familia": 40}
    assert snap["workspace"] == empty        # crear el espacio vacío no cambia el contenido del borrador


def test_assessments_are_read_when_present(db):
    opp = lic(db)
    lot = next(l for l in opp.lots if l.number == "1")
    req = db.scalars(select(Requirement).where(Requirement.code == "DEMO-REQ-001")).one()
    db.add(MatchAssessment(organization_id=opp.organization_id, opportunity_id=opp.id, lot_id=lot.id,
                           variant_id=lot.proposed_variant_id, requirement_id=req.id, result="coincide",
                           explanation="El catálogo declara titanio.", rule="eq", engine_version="1",
                           catalog_value="titanio", source_text="Material: titanio"))
    db.flush()
    snap = build_snapshot(db, opp)
    row = next(r for r in snap["requirements"] if r["code"] == "DEMO-REQ-001")
    assert row["assessments"][0]["result"] == "coincide"
    assert row["assessments"][0]["catalog_value"] == "titanio"
    other = next(r for r in snap["requirements"] if r["code"] == "DEMO-REQ-002")
    assert other["assessments"] == []


# ── F2 · plantillas DOCX ─────────────────────────────────────────────────────

from docx import Document as DocxDocument  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.services.generation.docx_templates import HEADER_TEXT, TEMPLATE_VERSION, TEMPLATES  # noqa: E402


def _paragraph_texts(element) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in element.iter(qn("w:p"))]


def docx_parts(data: bytes) -> dict:
    """Texto del cuerpo (incluye tablas), del encabezado y del pie de un DOCX."""
    doc = DocxDocument(io.BytesIO(data))
    section = doc.sections[0]
    return {
        "body": "\n".join(_paragraph_texts(doc.element.body)),
        "header": "\n".join(_paragraph_texts(section.header._element)),
        "footer": "\n".join(_paragraph_texts(section.footer._element)),
        "doc": doc,
    }


def render_all(snapshot: dict) -> dict[str, bytes]:
    return {key: render(snapshot) for key, render in TEMPLATES.items()}


def tables_of(data: bytes) -> list[list[list[str]]]:
    doc = DocxDocument(io.BytesIO(data))
    return [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]


SIX = ["resumen", "matriz", "faltantes", "indice", "propuesta_tecnica", "propuesta_economica"]


def test_template_registry():
    assert TEMPLATE_VERSION == "1.0"
    assert list(TEMPLATES) == SIX


def test_all_templates_have_watermark(db):
    snap = snapshot_of(db)
    for key, data in render_all(snap).items():
        parts = docx_parts(data)
        assert config.WATERMARK in parts["footer"], key
        assert snap["data_version"] in parts["footer"], key
        assert HEADER_TEXT == "Dipromedic · Borrador de expediente · DEMO"
        assert parts["header"].strip() == HEADER_TEXT, key
        assert "Borrador · no apto para presentación" in parts["body"], key
        # La fecha de generación está en el pie, en hora de Ciudad de México.
        assert re.search(r"\d{2}/\d{2}/\d{4} \d{2}:\d{2}", parts["footer"]), key


def test_economic_totals_text(db):
    snap = snapshot_of(db)
    parts = docx_parts(render_all(snap)["propuesta_economica"])
    assert "98,000.00" in parts["body"]
    assert "Total: 198,400.00 MXN" in parts["body"]
    assert snap["fiscal"]["text"] in parts["body"]
    [table] = tables_of(render_all(snap)["propuesta_economica"])
    assert table[0] == ["Partida", "Descripción", "Cantidad", "Unidad", "Precio unitario", "Moneda", "Subtotal",
                        "Origen del precio"]
    row1 = table[1]
    assert row1[0] == "1" and row1[2] == "40" and row1[3] == "pieza" and row1[4] == "2,450.00"
    assert row1[5] == "MXN" and row1[6] == "98,000.00" and row1[7].startswith("SIMULADO")


def test_economic_partial_total_shows_pending_prices(db):
    opp = lic(db)
    clear_lot_prices(db, opp, "2")
    next(l for l in opp.lots if l.number == "2").proposed_variant_id = real_variant(db, opp.organization_id).id
    db.flush()
    db.expire_all()
    snap = snapshot_of(db)
    data = render_all(snap)["propuesta_economica"]
    body = docx_parts(data)["body"]
    assert "Total parcial: 144,800.00 MXN" in body
    row2 = tables_of(data)[0][2]
    assert "[PENDIENTE: precio]" in row2[4] and "[PENDIENTE: precio]" in row2[6]
    assert "SIMULADO" not in row2[7]
    assert "[PENDIENTE: aprobación comercial de precios]" in body


def test_pending_markers_for_open_requirements(db):
    snap = snapshot_of(db)
    open_mandatory = [r for r in snap["requirements"] if r["open"] and r["mandatory"] is not False]
    assert len(open_mandatory) >= 5
    texts = {k: docx_parts(v)["body"] for k, v in render_all(snap).items()}
    for key in ("matriz", "faltantes"):
        assert texts[key].count("[PENDIENTE") >= len(open_mandatory), key
    # El código de cada requisito abierto aparece en la lista de faltantes.
    for req in open_mandatory:
        assert req["code"] in texts["faltantes"]


def test_faltantes_assigns_owner_and_due_date_or_marks_pending(db):
    snap = snapshot_of(db)
    [table, *_] = tables_of(render_all(snap)["faltantes"])
    header = table[0]
    assert header[:2] == ["No.", "Código"] and "Responsable" in header and "Fecha límite" in header
    by_code = {row[1]: dict(zip(header, row)) for row in table[1:]}
    task_reqs = {r["code"]: r["open_task"] for r in snap["requirements"] if r["open_task"]}
    assert task_reqs, "el escenario trae tareas abiertas"
    code, task = next(iter(task_reqs.items()))
    assert task["assignee"] in by_code[code]["Responsable"]
    no_task = next(r["code"] for r in snap["requirements"] if r["open"] and not r["open_task"] and r["code"] in by_code)
    assert "[PENDIENTE: responsable]" in by_code[no_task]["Responsable"]


FORBIDDEN_WORDS = re.compile(r"\b(firm\w*|sell\w*|presentad\w*)\b", re.IGNORECASE)


def test_forbidden_word_regex_matches_words_not_substrings():
    for bad in ("Firma", "firmado", "FIRMAS", "sello", "sellos", "Presentado", "presentada"):
        assert FORBIDDEN_WORDS.search(f"texto {bad} texto"), bad
    for fine in ("afirma", "no se afirma disponibilidad", "confirma ni descarta", "presentación", "información"):
        assert not FORBIDDEN_WORDS.search(f"texto {fine} texto"), fine


def _package_texts(snap: dict, data: bytes) -> dict[str, str]:
    """Todo el texto de un paquete: DOCX (cuerpo, encabezado y pie), CSV, valores del XLSX, README y manifiesto."""
    zf = open_zip(data)
    texts = {}
    for name in zf.namelist():
        content = zf.read(name)
        if name.endswith(".docx"):
            parts = docx_parts(content)
            texts[name] = "\n".join([parts["body"], parts["header"], parts["footer"]])
        elif name == "matriz.xlsx":
            wb = openpyxl.load_workbook(io.BytesIO(content))
            texts[name] = "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                                    if isinstance(c.value, str))
        else:
            texts[name] = content.decode("utf-8-sig")
    return texts


def test_no_signature_words(db):
    """Sobre lo que realmente se siembra: ni las plantillas ni los datos citados traen firma, sello o «presentado»."""
    codes = [c for (c,) in db.execute(select(Opportunity.external_id).order_by(Opportunity.external_id))]
    assert len([c for c in codes if c.startswith("DEMO-LIC-")]) >= 8
    seen_afirma = False
    for code in codes:
        snap = snapshot_of(db, code)
        data, _ = build_package(snap, render_all(snap))
        for name, text in _package_texts(snap, data).items():
            text = text.replace(config.WATERMARK, "")
            hit = FORBIDDEN_WORDS.search(text)
            assert not hit, (code, name, hit.group(0))
            # Tampoco hay espacios para rúbricas ni líneas de firma.
            assert "____" not in text, (code, name)
            seen_afirma = seen_afirma or "afirma" in text
    assert seen_afirma    # el texto de DEMO-REQ-006 («no se afirma…») sí está en la salida y no cuenta como coincidencia


def test_templates_render_with_minimal_data_and_mark_everything_pending(db):
    """Una oportunidad sin partidas, sin requisitos ni fechas: nada se inventa y todo queda pendiente."""
    org = db.scalars(select(Organization)).first()
    opp = Opportunity(organization_id=org.id, external_id="DEMO-VACIA", normalized_id="DEMOVACIA", title="Vacía",
                      buyer="Institución ficticia", data_kind="SYNTHETIC", verification_status="simulada")
    db.add(opp)
    db.flush()
    snap = build_snapshot(db, opp)
    assert snap["totals"]["lots_total"] == 0 and snap["totals"]["known_total"] is None
    texts = {k: docx_parts(v)["body"] for k, v in render_all(snap).items()}
    assert "[PENDIENTE: fecha de presentación]" in texts["resumen"]
    assert "[PENDIENTE: unidad compradora]" in texts["resumen"]
    assert "Afinidad no calculada" in texts["resumen"]
    assert "Sin evaluación técnica" in texts["propuesta_tecnica"] or "partidas" in texts["propuesta_tecnica"]
    assert "Total no calculable" in texts["propuesta_economica"]
    assert "No hay requisitos capturados" in texts["matriz"]


def test_technical_proposal_shows_both_values_and_never_grants_compliance(db):
    snap = snapshot_of(db)
    data = render_all(snap)["propuesta_tecnica"]
    body = docx_parts(data)["body"]
    assert "Sin evaluación técnica" in body
    assert "DEMO-PLA-003" in body and "DEMO-SIN-VALIDEZ" in body
    # Medida incompatible a la vista: la partida 2 exige 11 orificios y el catálogo declara 10.
    tables = tables_of(data)
    flat = [cell for t in tables for row in t for cell in row]
    assert any("11" in c for c in flat) and any(c.strip() == "10" or c.startswith("10 ") for c in flat)
    assert "Cumple ·" not in body and "cumple_validado" not in body


def test_engine_results_and_operators_are_spanish_in_drafts(db):
    """Integración 2: lo que el motor de E escribe (`coincide`, `incumple`, `desconocido`, `contains`) sale en
    español en la matriz y en la propuesta técnica, nunca el código crudo."""
    from app.models import User
    from app.services import workspace
    from app.services.generation.matrix import matrix_rows

    opp = lic(db)
    user = db.scalars(select(User).where(User.email == "analista1@dipromedic.example")).one()
    workspace.assess(db, user, opp)
    db.flush()
    snap = build_snapshot(db, opp)
    cells = [c for row in matrix_rows(snap) for c in row]
    text = "\n".join(cells)
    assert "Coincide: " in text and "Incumple: " in text and "Desconocido: " in text
    assert not re.search(r"(^|\n)(coincide|incumple|desconocido|no_comparable): ", text)
    assert "(contiene titanio)" in text and "contains" not in text
    body = docx_parts(render_all(snap)["propuesta_tecnica"])["body"]
    flat = " ".join(c for t in tables_of(render_all(snap)["propuesta_tecnica"]) for row in t for c in row)
    assert "Incumple: " in flat and "incumple: " not in flat and "contains" not in body + flat
    assert "contiene titanio" in flat


def test_technical_proposal_cites_assessments_and_catalog_credit(db):
    opp = lic(db)
    lot = next(l for l in opp.lots if l.number == "1")
    variant = real_variant(db, opp.organization_id)
    lot.proposed_variant_id = variant.id
    req = db.scalars(select(Requirement).where(Requirement.code == "DEMO-REQ-002")).one()
    db.add(MatchAssessment(organization_id=opp.organization_id, opportunity_id=opp.id, lot_id=lot.id,
                           variant_id=variant.id, requirement_id=req.id, result="coincide",
                           explanation="La ficha lista 6 orificios.", rule="eq", engine_version="1",
                           catalog_value="6", source_text="6 orificios"))
    db.flush()
    db.expire_all()
    snap = snapshot_of(db)
    body = docx_parts(render_all(snap)["propuesta_tecnica"])["body"]
    assert "La ficha lista 6 orificios." in body and "6 orificios" in body
    assert ("Contenido de catálogo: DIPROMEDIC S.A. DE C.V. · www.dipromedic.com · consultado el 2026-10-28"
            in body)
    assert "[PENDIENTE: precio]" not in body


def test_index_lists_documents_without_bundling_them(db):
    snap = snapshot_of(db)
    body = docx_parts(render_all(snap)["indice"])["body"]
    assert snap["documents"], "el escenario trae documentos"
    for doc in snap["documents"]:
        assert doc["code"] in body
    assert "No se incluye el archivo" in body
    assert "01_resumen_oportunidad.docx" in body and "manifest.json" in body
    assert "[PENDIENTE: documento" in body      # requisitos abiertos sin evidencia


def test_xml_control_characters_do_not_break_rendering(db):
    req = db.scalars(select(Requirement).where(Requirement.code == "DEMO-REQ-001")).one()
    req.text = "Placa\x1f de\x0b titanio"      # Postgres no admite NUL; el resto de los controles sí
    db.flush()
    db.expire_all()
    body = docx_parts(render_all(snapshot_of(db))["matriz"])["body"]
    assert "Placa de titanio" in body


# ── F3 · paquete ZIP con manifiesto ──────────────────────────────────────────

import csv  # noqa: E402

import openpyxl  # noqa: E402

from app.services.generation.package import build_package  # noqa: E402


def make_package(db, snapshot=None):
    snap = snapshot or snapshot_of(db)
    data, manifest = build_package(snap, render_all(snap))
    return snap, data, manifest


def open_zip(data: bytes) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(data))


def test_zip_layout(db):
    snap, data, manifest = make_package(db)
    names = open_zip(data).namelist()
    assert sorted(n for n in names if n.startswith("docx/")) == [
        "docx/01_resumen_oportunidad.docx", "docx/02_matriz_cumplimiento.docx", "docx/03_lista_faltantes.docx",
        "docx/04_indice_expediente.docx", "docx/05_propuesta_tecnica_preliminar.docx",
        "docx/06_propuesta_economica_demo.docx"]
    assert {"matriz.csv", "matriz.xlsx", "README_PENDIENTES.md", "manifest.json"} <= set(names)
    assert len(names) == 10 and len(set(names)) == 10


def test_manifest_hashes_match(db):
    snap, data, manifest = make_package(db)
    zf = open_zip(data)
    listed = {f["name"]: f for f in manifest["files"]}
    assert set(listed) == set(zf.namelist()) - {"manifest.json"}      # el manifiesto no se lista a sí mismo
    for name, entry in listed.items():
        content = zf.read(name)
        assert entry["sha256"] == hashlib.sha256(content).hexdigest(), name
        assert entry["size_bytes"] == len(content), name
    assert json.loads(zf.read("manifest.json").decode("utf-8")) == manifest
    assert manifest["template_version"] == TEMPLATE_VERSION
    assert manifest["generated_at"] == snap["generated_at"] and manifest["data_version"] == snap["data_version"]
    assert manifest["opportunity"] == {"external_id": "DEMO-LIC-001", "data_kind": "SYNTHETIC"}
    assert manifest["warnings"] == snap["warnings"] and manifest["notice"] == config.WATERMARK
    assert set(manifest["provenance"]) == {"REAL", "SYNTHETIC", "HISTORICAL"}
    assert manifest["provenance"]["SYNTHETIC"] > 0 and manifest["provenance"]["REAL"] == 0


def test_zip_is_reproducible_from_the_same_inputs(db):
    snap = snapshot_of(db)
    first, _ = build_package(snap, render_all(snap))
    time.sleep(1.1)       # el reloj de pared no debe influir: ni en los DOCX ni en el XLSX ni en el ZIP
    again, _ = build_package(snap, render_all(snap))
    za, zb = open_zip(first), open_zip(again)
    diff = [n for n in za.namelist() if za.read(n) != zb.read(n)]
    assert not diff, diff
    assert [(i.filename, i.date_time, i.compress_type, i.external_attr) for i in za.infolist()] == \
        [(i.filename, i.date_time, i.compress_type, i.external_attr) for i in zb.infolist()]
    assert hashlib.sha256(first).hexdigest() == hashlib.sha256(again).hexdigest()
    zf = open_zip(first)
    assert len({i.date_time for i in zf.infolist()}) == 1
    for name in ("matriz.xlsx", "docx/06_propuesta_economica_demo.docx"):
        inner = zipfile.ZipFile(io.BytesIO(zf.read(name)))
        assert {i.date_time for i in inner.infolist()} == {zf.infolist()[0].date_time}


def test_provenance_counts_real_records(db):
    opp = lic(db)
    next(l for l in opp.lots if l.number == "2").proposed_variant_id = real_variant(db, opp.organization_id).id
    db.flush()
    db.expire_all()
    _, _, manifest = make_package(db)
    assert manifest["provenance"]["REAL"] == 1       # la variante REAL
    assert manifest["provenance_detail"]["variants"] == {"REAL": 1, "SYNTHETIC": 2}


def test_csv_utf8_bom_and_one_row_per_requirement(db):
    snap, data, _ = make_package(db)
    raw = open_zip(data).read("matriz.csv")
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Código", "Partida", "Requisito"]
    assert len(rows) - 1 == len(snap["requirements"])
    assert all(r[-1] == config.WATERMARK for r in rows[1:])
    by_code = {r[0]: r for r in rows[1:]}
    assert "[PENDIENTE: evidencia vinculada]" in by_code["DEMO-REQ-005"]


def test_xlsx_money_cells_are_numbers_with_format_and_the_watermark(db):
    """El XLSX guarda montos como números de hoja de cálculo (coma flotante), no como texto exacto."""
    snap, data, _ = make_package(db)
    wb = openpyxl.load_workbook(io.BytesIO(open_zip(data).read("matriz.xlsx")))
    assert wb.sheetnames == ["Resumen", "Matriz", "Partidas"]
    for ws in wb.worksheets:
        assert ws["A1"].value == config.WATERMARK, ws.title
    lots = wb["Partidas"]
    header = [c.value for c in lots[3]]
    assert header[:8] == ["Partida", "Descripción", "Cantidad", "Unidad", "Precio unitario", "Moneda", "Subtotal",
                          "Origen del precio"]
    row1 = [c.value for c in lots[4]]
    assert row1[0] == "1" and Decimal(str(row1[2])) == Decimal("40") and Decimal(str(row1[4])) == Decimal("2450")
    assert Decimal(str(row1[6])) == Decimal("98000.00") and row1[7].startswith("SIMULADO")
    assert isinstance(row1[6], (int, float))
    assert lots["G4"].number_format == "#,##0.00" and lots["E4"].number_format == "#,##0.00"
    notes = [c.value for row in lots.iter_rows(min_row=8) for c in row if isinstance(c.value, str)]
    assert any("coma flotante" in n and "DOCX" in n for n in notes)
    summary = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=3) if r[0].value}
    assert summary["Versión de datos"] == snap["data_version"]
    assert summary["Total"] == "Total: 198,400.00 MXN"


def test_xlsx_money_round_trips_to_the_exact_decimal_on_awkward_cents(db):
    opp = lic(db)
    set_lot_price(db, opp, "1", "0.07")
    set_lot_price(db, opp, "2", "99999999999.99")
    lot3 = set_lot_price(db, opp, "3", "33.33")
    lot3.quantity = Decimal("2.500")            # 2.5 × 33.33 = 83.325 → 83.33 (mitad hacia arriba)
    db.flush()
    db.expire_all()
    snap, data, _ = make_package(db)
    assert [l["subtotal"] for l in snap["lots"]] == ["2.80", "1999999999999.80", "83.33"]
    lots = openpyxl.load_workbook(io.BytesIO(open_zip(data).read("matriz.xlsx")))["Partidas"]
    rows = [[c.value for c in lots[r]] for r in (4, 5, 6)]
    for lot, row in zip(snap["lots"], rows):
        assert Decimal(str(row[2])) == Decimal(lot["quantity"]), lot["number"]
        assert Decimal(str(row[4])) == Decimal(lot["unit_price"]), lot["number"]
        assert Decimal(str(row[6])) == Decimal(lot["subtotal"]), lot["number"]
    assert Decimal(str(rows[0][4])) == Decimal("0.07") and Decimal(str(rows[1][4])) == Decimal("99999999999.99")
    # El texto exacto sigue siendo el de los DOCX.
    body = docx_parts(render_all(snap)["propuesta_economica"])["body"]
    assert "1,999,999,999,999.80" in body and "99,999,999,999.99" in body
    assert "Total: 2,000,000,000,085.93 MXN" in body      # 2.80 + 1,999,999,999,999.80 + 83.33, sumado con Decimal


def test_xlsx_integral_quantities_have_no_trailing_point(db):
    opp = lic(db)
    next(l for l in opp.lots if l.number == "3").quantity = Decimal("2.500")
    db.flush()
    db.expire_all()
    _, data, _ = make_package(db)
    lots = openpyxl.load_workbook(io.BytesIO(open_zip(data).read("matriz.xlsx")))["Partidas"]
    assert lots["C4"].number_format == "#,##0" and lots["C5"].number_format == "#,##0"      # 40 y 20, no «40.»
    assert lots["C6"].number_format == "#,##0.###"                                          # 2.5


def test_spreadsheet_formula_injection_is_neutralised(db):
    req = db.scalars(select(Requirement).where(Requirement.code == "DEMO-REQ-001")).one()
    req.text = '=HYPERLINK("http://x","clic")'
    db.flush()
    db.expire_all()
    _, data, _ = make_package(db)
    zf = open_zip(data)
    rows = list(csv.reader(io.StringIO(zf.read("matriz.csv").decode("utf-8-sig"))))
    assert next(r for r in rows if r[0] == "DEMO-REQ-001")[2].startswith("'=HYPERLINK(\"http://x\",\"clic\")")
    wb = openpyxl.load_workbook(io.BytesIO(zf.read("matriz.xlsx")))
    cell = next(r[2] for r in wb["Matriz"].iter_rows(min_row=4) if r[0].value == "DEMO-REQ-001")
    assert cell.data_type == "s" and cell.value.startswith("'=")


def test_readme_lists_pending(db):
    snap, data, _ = make_package(db)
    readme = open_zip(data).read("README_PENDIENTES.md").decode("utf-8")
    assert config.WATERMARK in readme and "Borrador · no apto para presentación" in readme
    assert snap["data_version"] in readme
    for req in snap["requirements"]:
        if req["open"]:
            assert req["code"] in readme
    for warning in snap["warnings"]:
        assert warning in readme
    assert "manifest.json" in readme
    assert not re.search(r"\b(firma\w*|sello\w*|presentad[oa]s?)\b", readme.replace(config.WATERMARK, ""), re.I)


def test_third_party_excluded(db):
    opp = lic(db)
    secret = b"%PDF-1.4 contenido de un tercero que no debe viajar en el ZIP"
    third = Document(organization_id=opp.organization_id, code="TER-001", title="Certificado de un tercero",
                     scope="expediente", data_kind="REAL", opportunity_id=opp.id, third_party=True,
                     external_url="https://tercero.example/cert.pdf", review_status="pendiente")
    db.add(third)
    db.flush()
    db.expire_all()
    snap, data, manifest = make_package(db)
    zf = open_zip(data)
    entry = next(d for d in manifest["excluded_documents"] if d["code"] == "TER-001")
    assert entry["third_party"] is True and "tercero" in entry["reason"].lower()
    assert entry["external_url"] == "https://tercero.example/cert.pdf"
    # Ningún archivo de documentos viaja en el paquete: solo plantillas, matriz, README y manifiesto.
    assert all(n.startswith("docx/") or n in {"matriz.csv", "matriz.xlsx", "README_PENDIENTES.md", "manifest.json"}
               for n in zf.namelist())
    assert not any(secret in zf.read(n) for n in zf.namelist())
    # También se lista el resto de los documentos del expediente, con su motivo.
    assert {d["code"] for d in manifest["excluded_documents"]} >= {d["code"] for d in snap["documents"]}
    assert all(d["reason"] for d in manifest["excluded_documents"])


# ── F4 · trabajo asíncrono y endpoints ───────────────────────────────────────

import uuid  # noqa: E402

from app.core.security import get_password_hash  # noqa: E402
from app.models import Job, User, UserRole  # noqa: E402
from app.services import jobs  # noqa: E402
from app.services.generation import artifact_id_for  # noqa: E402
from app.services.storage import get_storage  # noqa: E402
from tests.conftest import login  # noqa: E402


@pytest.fixture
def responsable(client):
    return login(client, "responsable@dipromedic.example")


def opp_id(code: str = "DEMO-LIC-001") -> str:
    with SessionLocal() as db:
        return lic(db, code).id


def run_jobs() -> None:
    """El worker está apagado en las pruebas: se ejecuta lo encolado, como lo haría el hilo del servicio."""
    with SessionLocal() as db:
        while jobs.run_pending_once(db) is not None:
            pass


def generate_now(client, headers, templates=None, key=None, code="DEMO-LIC-001") -> dict:
    """POST generar + ejecutar el trabajo + devolver el trabajo terminado."""
    body = {"templates": templates} if templates is not None else None
    r = client.post(f"/api/v1/workspaces/{opp_id(code)}/generate", headers={**headers, "Idempotency-Key": key or uuid.uuid4().hex},
                    json=body)
    assert r.status_code == 202, r.text
    run_jobs()
    job = client.get(f"/api/v1/jobs/{r.json()['job_id']}", headers=headers).json()
    assert job["status"] == "done", job
    return job


def artifacts_of(client, headers, code="DEMO-LIC-001") -> list[dict]:
    r = client.get(f"/api/v1/workspaces/{opp_id(code)}/artifacts", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def test_generate_idempotent(client, analista):
    url = f"/api/v1/workspaces/{opp_id()}/generate"
    headers = {**analista, "Idempotency-Key": "doble-clic-1"}
    first = client.post(url, headers=headers)
    again = client.post(url, headers=headers)
    other = client.post(url, headers={**analista, "Idempotency-Key": "doble-clic-2"})
    assert first.status_code == again.status_code == other.status_code == 202
    assert first.json()["job_id"] == again.json()["job_id"] != other.json()["job_id"]
    with SessionLocal() as db:
        assert db.query(Job).filter(Job.kind == "generate_artifact").count() == 2
    run_jobs()
    # Repetir la clave con el trabajo ya terminado sigue devolviendo ese mismo trabajo y no crea otro paquete.
    done = client.post(url, headers=headers)
    assert done.json()["job_id"] == first.json()["job_id"] and done.json()["status"] == "done"
    assert len(artifacts_of(client, analista)) == 2


def test_generate_runs_as_a_job_with_progress_and_audit(client, analista):
    job = generate_now(client, analista)
    assert job["kind"] == "generate_artifact" and job["progress"] == 100 and job["message"] == "Borrador listo"
    result = job["result"]
    [artifact] = artifacts_of(client, analista)
    assert result["artifact_id"] == artifact["id"] == artifact_id_for(job["id"])
    assert artifact["status"] == "borrador" and artifact["sha256"] == result["sha256"]
    assert re.fullmatch(r"[0-9a-f]{64}", artifact["data_version"]) and artifact["template_version"] == "1.0"
    assert artifact["created_by"]["email"] == "analista1@dipromedic.example"
    assert artifact["filename"].startswith("borrador_expediente_DEMO-LIC-001_") and artifact["filename"].endswith(".zip")
    assert any("Afinidad no calculada" in w for w in artifact["warnings"])
    with SessionLocal() as db:
        org_id = db.scalars(select(Organization.id)).first()
        row = db.get(GeneratedArtifact, artifact["id"])
        assert row.status == "borrador" and row.storage_key == f"artifacts/{org_id}/{opp_id()}/{artifact['id']}.zip"
        data = get_storage().get(row.storage_key)
        assert hashlib.sha256(data).hexdigest() == artifact["sha256"] and len(data) == artifact["size_bytes"]
        assert row.snapshot["data_version"] == artifact["data_version"]
        ws = db.scalars(select(BidWorkspace).where(BidWorkspace.opportunity_id == opp_id())).one()
        assert row.workspace_id == ws.id
        events = [e for e in db.scalars(select(AuditEvent).where(AuditEvent.action.like("artifact.%")))]
        [gen] = [e for e in events if e.action == "artifact.generate"]
        assert gen.resource_id == artifact["id"] and gen.actor.email == "analista1@dipromedic.example"
        assert gen.after["sha256"] == artifact["sha256"]
        assert audit.verify_chain(db, org_id)[0] is True


def test_generate_reuses_existing_workspace(client, analista):
    with SessionLocal() as db:
        opp = lic(db)
        ws = BidWorkspace(organization_id=opp.organization_id, opportunity_id=opp.id, decision="participar",
                          affinity={"components": {"familia": 40}, "total": 75, "coverage_complete": True,
                                    "blocking": []})
        db.add(ws)
        db.commit()
        ws_id = ws.id
    generate_now(client, analista)
    generate_now(client, analista)
    with SessionLocal() as db:
        assert db.query(BidWorkspace).count() == 1
        assert {a.workspace_id for a in db.query(GeneratedArtifact)} == {ws_id}
        # Lo que E guardó en el espacio de trabajo viaja en la instantánea del borrador.
        assert db.query(GeneratedArtifact).first().snapshot["affinity"]["total"] == 75


def test_generated_package_is_readable_end_to_end(client, analista):
    generate_now(client, analista)
    [artifact] = artifacts_of(client, analista)
    r = client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=analista)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"
    assert r.headers["content-disposition"].startswith("attachment; filename=\"borrador_expediente_DEMO-LIC-001_")
    assert r.headers["cache-control"] == "private, no-store" and r.headers["x-content-type-options"] == "nosniff"
    assert hashlib.sha256(r.content).hexdigest() == artifact["sha256"]
    zf = open_zip(r.content)
    manifest = json.loads(zf.read("manifest.json"))
    assert manifest["data_version"] == artifact["data_version"]
    for entry in manifest["files"]:
        assert hashlib.sha256(zf.read(entry["name"])).hexdigest() == entry["sha256"]
    body = docx_parts(zf.read("docx/06_propuesta_economica_demo.docx"))["body"]
    assert "Total: 198,400.00 MXN" in body
    # La vista del manifiesto desde la UI devuelve lo mismo que el archivo.
    m = client.get(f"/api/v1/artifacts/{artifact['id']}/manifest", headers=analista)
    assert m.status_code == 200 and m.json() == manifest


def test_generate_with_template_subset(client, analista):
    generate_now(client, analista, templates=["indice", "resumen", "resumen"])
    [artifact] = artifacts_of(client, analista)
    zf = open_zip(client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=analista).content)
    docx_names = sorted(n for n in zf.namelist() if n.startswith("docx/"))
    assert docx_names == ["docx/01_resumen_oportunidad.docx", "docx/04_indice_expediente.docx"]
    # El índice describe lo que de verdad viaja en este paquete.
    index = docx_parts(zf.read("docx/04_indice_expediente.docx"))["body"]
    assert "01_resumen_oportunidad.docx" in index and "06_propuesta_economica_demo.docx" not in index


def test_generate_rejects_unknown_or_empty_templates(client, analista):
    url = f"/api/v1/workspaces/{opp_id()}/generate"
    bad = client.post(url, headers=analista, json={"templates": ["resumen", "carta_firmada"]})
    assert bad.status_code == 422 and bad.json()["code"] == "validation_error"
    assert bad.json()["details"][0]["field"] == "templates"
    assert client.post(url, headers=analista, json={"templates": []}).status_code == 422
    assert client.post(url, headers={**analista, "Idempotency-Key": "x" * 201}).status_code == 422
    with SessionLocal() as db:
        assert db.query(Job).count() == 0


def test_download_other_org_404(client, analista):
    generate_now(client, analista)
    [artifact] = artifacts_of(client, analista)
    with SessionLocal() as db:
        other = Organization(name="Otra organización")
        db.add(other)
        db.flush()
        user = User(organization_id=other.id, email="intruso@otra.example", full_name="Intruso",
                    hashed_password=get_password_hash("demo-pass"))
        user.roles = [UserRole(role="analista_comercial")]
        db.add(user)
        db.commit()
    intruder = login(client, "intruso@otra.example")
    assert client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=intruder).status_code == 404
    assert client.get(f"/api/v1/artifacts/{artifact['id']}/manifest", headers=intruder).status_code == 404
    assert client.get(f"/api/v1/workspaces/{opp_id()}/artifacts", headers=intruder).status_code == 404
    assert client.post(f"/api/v1/workspaces/{opp_id()}/generate", headers=intruder).status_code == 404
    assert client.get(f"/api/v1/workspaces/{opp_id()}/generation-preview", headers=intruder).status_code == 404
    with SessionLocal() as db:
        assert db.query(AuditEvent).filter(AuditEvent.action == "artifact.download").count() == 0


def test_admin_ti_cannot_download(client, analista, ti):
    generate_now(client, analista)
    [artifact] = artifacts_of(client, analista)
    r = client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=ti)
    assert r.status_code == 403 and r.json()["code"] == "forbidden"
    assert client.get(f"/api/v1/artifacts/{artifact['id']}/manifest", headers=ti).status_code == 403
    assert client.get(f"/api/v1/workspaces/{opp_id()}/artifacts", headers=ti).status_code == 403
    assert client.post(f"/api/v1/workspaces/{opp_id()}/generate", headers=ti).status_code == 403
    with SessionLocal() as db:
        assert db.query(AuditEvent).filter(AuditEvent.action == "artifact.download").count() == 0


def test_direccion_downloads_but_does_not_generate(client, analista, direccion):
    generate_now(client, analista)
    [artifact] = artifacts_of(client, direccion)
    assert client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=direccion).status_code == 200
    assert client.post(f"/api/v1/workspaces/{opp_id()}/generate", headers=direccion).status_code == 403


def test_regenerate_keeps_previous_hash(client, analista):
    first_job = generate_now(client, analista, key="k1")
    [first] = artifacts_of(client, analista)
    with SessionLocal() as db:
        opp = lic(db)
        set_lot_price(db, opp, "1", "1600.00")        # cambian los datos entre una generación y otra
        db.commit()
    second_job = generate_now(client, analista, key="k2")
    rows = artifacts_of(client, analista)
    assert len(rows) == 2 and rows[0]["id"] == second_job["result"]["artifact_id"]      # la más reciente primero
    old = next(r for r in rows if r["id"] == first["id"])
    new = rows[0]
    assert old["sha256"] == first["sha256"] and old["data_version"] == first["data_version"]
    assert new["data_version"] != old["data_version"] and new["sha256"] != old["sha256"]
    # El ZIP anterior sigue intacto en el almacenamiento y coincide con la huella que se registró.
    got = client.get(f"/api/v1/artifacts/{old['id']}/download", headers=analista).content
    assert hashlib.sha256(got).hexdigest() == first["sha256"]
    assert first_job["result"]["artifact_id"] == first["id"]
    assert {a["status"] for a in rows} == {"borrador"}
    # Y el nuevo ya trae el precio de la partida, no el del catálogo.
    body = docx_parts(open_zip(client.get(f"/api/v1/artifacts/{new['id']}/download", headers=analista).content)
                      .read("docx/06_propuesta_economica_demo.docx"))["body"]
    assert "64,000.00" in body and "1,600.00" in body


def test_download_audited(client, analista, direccion):
    generate_now(client, analista)
    [artifact] = artifacts_of(client, analista)
    assert client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=direccion).status_code == 200
    with SessionLocal() as db:
        [event] = db.scalars(select(AuditEvent).where(AuditEvent.action == "artifact.download"))
        assert event.resource_type == "generated_artifact" and event.resource_id == artifact["id"]
        assert event.actor.email == "direccion@dipromedic.example"
        assert event.after["sha256"] == artifact["sha256"] and event.after["filename"] == artifact["filename"]
        assert audit.verify_chain(db, event.organization_id)[0] is True


def test_corrupted_package_is_not_delivered(client, analista):
    generate_now(client, analista)
    [artifact] = artifacts_of(client, analista)
    with SessionLocal() as db:
        key = db.get(GeneratedArtifact, artifact["id"]).storage_key
    storage = get_storage()
    storage.put(key, storage.get(key) + b"x", "application/zip")
    r = client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=analista)
    assert r.status_code == 409 and r.json()["code"] == "integrity_error"
    storage.delete(key)
    r = client.get(f"/api/v1/artifacts/{artifact['id']}/download", headers=analista)
    assert r.status_code == 404 and r.json()["code"] == "file_missing"
    with SessionLocal() as db:
        assert db.query(AuditEvent).filter(AuditEvent.action == "artifact.download").count() == 0


def test_failed_generation_leaves_no_artifact_or_orphan_file(client, analista, monkeypatch):
    from app.services import generation

    with SessionLocal() as db:          # el espacio de trabajo ya existe: la falla ocurre después de guardar el ZIP
        opp = lic(db)
        db.add(BidWorkspace(organization_id=opp.organization_id, opportunity_id=opp.id))
        db.commit()

    def boom(*args, **kwargs):
        raise RuntimeError("falla de auditoría simulada")

    monkeypatch.setattr(generation.audit, "record", boom)
    r = client.post(f"/api/v1/workspaces/{opp_id()}/generate", headers={**analista, "Idempotency-Key": "falla-1"})
    run_jobs()
    job = client.get(f"/api/v1/jobs/{r.json()['job_id']}", headers=analista).json()
    assert job["status"] == "failed" and "falla de auditoría simulada" in job["error"]
    assert artifacts_of(client, analista) == []
    with SessionLocal() as db:
        org_id = db.scalars(select(Organization.id)).first()
    key = f"artifacts/{org_id}/{opp_id()}/{artifact_id_for(job['id'])}.zip"
    assert not get_storage().exists(key)
    # Repetir la misma clave devuelve el mismo trabajo fallido; la UI usa una clave nueva en cada clic.
    again = client.post(f"/api/v1/workspaces/{opp_id()}/generate", headers={**analista, "Idempotency-Key": "falla-1"})
    assert again.json()["job_id"] == r.json()["job_id"] and again.json()["status"] == "failed"


def test_preview_matches_what_gets_generated(client, analista, direccion):
    url = f"/api/v1/workspaces/{opp_id()}/generation-preview"
    before = client.get(url, headers=direccion)
    assert before.status_code == 200, before.text
    preview = before.json()
    assert preview["notice"] == config.WATERMARK and preview["changed_since_last"] is False
    assert preview["last_data_version"] is None
    assert [t["key"] for t in preview["templates"]] == SIX
    assert len(preview["open_requirements"]) == 7 and preview["totals_text"] == "Total: 198,400.00 MXN"
    assert any("Afinidad no calculada" in w for w in preview["warnings"])
    with SessionLocal() as db:
        assert db.query(GeneratedArtifact).count() == 0 and db.query(BidWorkspace).count() == 0   # solo lee
    generate_now(client, analista)
    [artifact] = artifacts_of(client, analista)
    after = client.get(url, headers=direccion).json()
    assert after["data_version"] == artifact["data_version"] == preview["data_version"]
    assert after["last_data_version"] == artifact["data_version"] and after["changed_since_last"] is False
    with SessionLocal() as db:
        set_lot_price(db, lic(db), "1", "1700.00")
        db.commit()
    changed = client.get(url, headers=direccion).json()
    assert changed["changed_since_last"] is True and changed["data_version"] != artifact["data_version"]
